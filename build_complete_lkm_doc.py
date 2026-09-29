# -*- coding: utf-8 -*-
"""
build_complete_lkm_doc.py
Generates LKM_Mengelola_Data_List_Koleksi_Flutter_Dart_LENGKAP.docx
by updating the student's partially filled document with complete, accurate answers,
removing trailing dots, fixing heading styles, and embedding real screenshots.
"""

import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

INPUT_DOCX = r"C:\Users\ASUS\Downloads\LKM_Mengelola_Data_List_Koleksi_Flutter_Dart.docx"
OUTPUT_DOCX = r"c:\SEMESTER 5\MOBILE PROGRAMMING\tugas 2\LKM_Mengelola_Data_List_Koleksi_Flutter_Dart_LENGKAP.docx"
SCREENSHOT_DIR = r"c:\SEMESTER 5\MOBILE PROGRAMMING\tugas 2\bukti_screenshot"

def get_img(name):
    p = os.path.join(SCREENSHOT_DIR, name)
    if not os.path.exists(p):
        print(f"WARNING: Image not found: {p}")
    return p

def remove_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_table_borders(table, color_hex="D0D7DE", sz="4"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color_hex)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def insert_p_after(ref_p, text, doc, style='Normal', space_before=2, space_after=4, bold_prefix=None):
    new_p = doc.add_paragraph(style=style)
    new_p.paragraph_format.space_before = Pt(space_before)
    new_p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_b = new_p.add_run(bold_prefix)
        r_b.bold = True
    new_p.add_run(text)
    ref_p._p.addnext(new_p._p)
    return new_p

def insert_code_after(ref_p, code_str, doc, space_before=3, space_after=4):
    new_p = doc.add_paragraph(style='CodeBlock')
    new_p.paragraph_format.space_before = Pt(space_before)
    new_p.paragraph_format.space_after = Pt(space_after)
    r = new_p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    ref_p._p.addnext(new_p._p)
    return new_p

def insert_callout_after(ref_p, doc, text, title=None, border_color="2A4D69", bg_color="F5F7FA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_color)
    tcBorders.append(left)
    for b_name in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), bg_color)
    tcPr.append(shd)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    if title:
        r_t = p.add_run(f"{title}\n")
        r_t.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = RGBColor(42, 77, 105)
    r = p.add_run(text)
    r.font.size = Pt(9)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

    ref_p._p.addnext(tbl._tbl)
    tbl._tbl.addnext(p_spacer._p)
    return p_spacer

def insert_image_after(ref_p, doc, img_path, caption, width=Inches(5.4)):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    if os.path.exists(img_path):
        p_img.add_run().add_picture(img_path, width=width)
    else:
        p_img.add_run(f"[Gambar tidak ditemukan: {img_path}]")

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption)
    r_cap.font.italic = True
    r_cap.font.size = Pt(8.5)
    r_cap.font.color.rgb = RGBColor(70, 70, 70)

    ref_p._p.addnext(p_img._p)
    p_img._p.addnext(p_cap._p)
    return p_cap

def insert_side_by_side_after(ref_p, doc, img1_path, cap1, img2_path, cap2, img_width=Inches(2.65)):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_borders(tbl)

    c1 = tbl.cell(0, 0)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(2)
    if os.path.exists(img1_path):
        p1.add_run().add_picture(img1_path, width=img_width)
    p_cap1 = c1.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap1.paragraph_format.space_before = Pt(2)
    p_cap1.paragraph_format.space_after = Pt(4)
    r1 = p_cap1.add_run(cap1)
    r1.font.italic = True
    r1.font.size = Pt(8)
    r1.font.color.rgb = RGBColor(70, 70, 70)

    c2 = tbl.cell(0, 1)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(2)
    if os.path.exists(img2_path):
        p2.add_run().add_picture(img2_path, width=img_width)
    p_cap2 = c2.add_paragraph()
    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap2.paragraph_format.space_before = Pt(2)
    p_cap2.paragraph_format.space_after = Pt(4)
    r2 = p_cap2.add_run(cap2)
    r2.font.italic = True
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor(70, 70, 70)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(6)

    ref_p._p.addnext(tbl._tbl)
    tbl._tbl.addnext(p_spacer._p)
    return p_spacer

def insert_table_after(ref_p, doc, headers, rows_data, col_widths=None, header_bg="2A4D69"):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color_hex="D0D7DE", sz="4")

    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        set_cell_background(cell, header_bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(rows_data):
        row_bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            set_cell_background(cell, row_bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if c_idx == 0 or len(str(val)) <= 10:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)

    if col_widths:
        for r in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = w

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(6)

    ref_p._p.addnext(tbl._tbl)
    tbl._tbl.addnext(p_spacer._p)
    return p_spacer

def main():
    print(f"Loading base document: {INPUT_DOCX}")
    doc = docx.Document(INPUT_DOCX)

    # 1. Update Student Table (Table 1)
    t1 = doc.tables[1]
    t1.cell(0, 1).text = "[Silakan isi Nama Lengkap Anda]"
    t1.cell(1, 1).text = "[Silakan isi NIM Anda]"
    t1.cell(2, 1).text = "[Silakan isi Kelas Anda]"
    t1.cell(3, 1).text = "[Silakan isi Tanggal Pengerjaan]"
    for r in t1.rows:
        for c in r.cells:
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.size = Pt(9.5)
    print("Table 1 student identity placeholders updated.")

    # 2. Identify and Clean/Update Paragraphs
    # We will build a map of paragraphs by index or key anchor text
    paras = list(doc.paragraphs)

    # Remove all dotted line paragraphs first to avoid clutter
    dots_removed = 0
    for p in paras:
        text = p.text.strip()
        if re.match(r'^\.{3,}$', text) or re.match(r'^[.…\s]{4,}$', text):
            p._p.getparent().remove(p._p)
            dots_removed += 1
    print(f"Removed {dots_removed} dotted line placeholder paragraphs.")

    # Refresh paragraph list after removal
    paras = list(doc.paragraphs)
    p_map = {p.text.strip(): p for p in paras}

    def find_p(prefix):
        for p in doc.paragraphs:
            if p.text.strip().startswith(prefix):
                return p
        return None

    # --- PRAKTIKUM 1 ---
    # Fix P040 from Heading 2 to Normal
    p_p1_heading = find_p("Praktikum 1")
    p_p1_out_lbl = find_p("Tuliskan hasil keluaran program:")
    if p_p1_out_lbl:
        # The explanation paragraph is after p_p1_out_lbl
        # Let's find the paragraph that starts with "Program menginisialisasi"
        p_expl = find_p("Program menginisialisasi")
        if p_expl:
            p_expl.style = 'Normal'
            p_expl.paragraph_format.space_before = Pt(2)
            p_expl.paragraph_format.space_after = Pt(4)
            print("Praktikum 1 explanation style changed from Heading 2 to Normal.")

    # --- PRAKTIKUM 2 ---
    # Prompt: jelaskan sederhana bahwa mahasiswa[0] mengambil elemen pada indeks 0 dan mahasiswa.first mengambil elemen pertama.
    # Hapus klaim “memory offset” dan klaim bahwa operator [] hanya ada pada List.
    p_p2_q = find_p("Apa perbedaan mahasiswa[0] dengan mahasiswa.first?")
    if p_p2_q:
        # Remove old confusing paragraphs until Praktikum 3
        curr = p_p2_q._p.getnext()
        while curr is not None:
            tag = curr.tag.split('}')[-1]
            text = "".join(curr.itertext()).strip()
            if "Praktikum 3" in text:
                break
            nxt = curr.getnext()
            curr.getparent().remove(curr)
            curr = nxt

        # Insert simple, correct answer
        cur_anchor = p_p2_q
        p_ans_out = insert_p_after(cur_anchor, "Ahmad\nAhmad\nDewi\n4", doc, style='CodeBlock', space_before=2, space_after=4)

        p_ans_expl = insert_p_after(p_ans_out,
            "Penjelasan Perbedaan mahasiswa[0] dan mahasiswa.first:\n"
            "1. Cara Kerja dan Mekanisme Akses:\n"
            "   • mahasiswa[0] mengambil elemen pada indeks tertentu (dalam hal ini indeks ke-0) menggunakan operator pengindeksan [] (subscript operator). Perlu dipahami bahwa operator [] ini bukan hanya milik List, melainkan juga digunakan pada struktur koleksi lain di Dart seperti Map.\n"
            "   • mahasiswa.first mengambil elemen pertama dari koleksi melalui properti getter bawaan dari interface Iterable (yang diimplementasikan oleh kelas List).\n"
            "2. Perilaku Saat List Kosong:\n"
            "   • Jika list dalam kondisi kosong (length = 0), pemanggilan mahasiswa[0] melempar exception: RangeError (Index out of range).\n"
            "   • Sementara itu, pemanggilan mahasiswa.first melempar exception: StateError (Bad state: No element).\n"
            "3. Fleksibilitas Akses:\n"
            "   • Operator [] dapat mengakses elemen di sembarang posisi indeks acak (misal: indeks 0, 1, 2, dst.), sedangkan properti getter .first secara khusus selalu hanya menunjuk pada elemen paling awal.",
            doc, style='Normal', space_before=2, space_after=6
        )
        print("Praktikum 2 updated with accurate, simplified explanation.")

    # --- PRAKTIKUM 3 & 4 ---
    p_p4_h = find_p("Praktikum 4")
    # Find the explanation paragraph of Praktikum 4
    p_p4_expl = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Alasan indeks"):
            p_p4_expl = p
            break

    if p_p4_expl:
        insert_image_after(p_p4_expl, doc, get_img("terminal_praktikum_1_4.png"),
            "Gambar D.1: Bukti Eksekusi Terminal dart run bin/praktikum_dasar.dart (Praktikum 1–4)", width=Inches(5.4))
        print("Screenshot terminal_praktikum_1_4.png inserted after Praktikum 4.")

    # --- PRAKTIKUM 5 ---
    p_p5_h = find_p("Praktikum 5")
    # Add detailed trace for Praktikum 5
    if p_p5_h:
        # find the codeblock after Praktikum 5
        curr = p_p5_h._p.getnext()
        # let's locate the codeblock
        code_p = None
        for p in doc.paragraphs:
            if p.text.strip().startswith("mahasiswa.remove('Budi Santoso');"):
                code_p = p
                break
        if code_p:
            p_p5_out = insert_p_after(code_p, "[]", doc, style='CodeBlock', space_before=2, space_after=4)
            insert_p_after(p_p5_out,
                "Penjelasan Tahapan Penghapusan Data:\n"
                "1. mahasiswa.remove('Budi Santoso'): Mencari objek string 'Budi Santoso' dan menghapus kemunculan pertamanya dari list. Data list menjadi: [Ahmad, Citra, Dewi, Citra, Dewi, Eka, Farhan].\n"
                "2. mahasiswa.removeAt(0): Menghapus elemen pada posisi indeks 0 (yaitu 'Ahmad'). Seluruh elemen berikutnya otomatis bergeser satu posisi ke kiri. Data list menjadi: [Citra, Dewi, Citra, Dewi, Eka, Farhan].\n"
                "3. mahasiswa.removeLast(): Menghapus elemen yang berada di posisi paling akhir list (yaitu 'Farhan'). Data list menjadi: [Citra, Dewi, Citra, Dewi, Eka].\n"
                "4. mahasiswa.clear(): Mengosongkan seluruh elemen yang masih ada di dalam list secara instan. Data list menjadi kosong: [] dengan panjang length = 0.",
                doc, style='Normal', space_before=2, space_after=6
            )
            print("Praktikum 5 detailed trace added.")

    # --- PRAKTIKUM 6 ---
    # Prompt: jelaskan bahwa setelah clear() di Praktikum 5, iterasi pada List kosong tidak mencetak nama.
    # Jika List diisi ulang untuk demonstrasi, tandai sebagai langkah tambahan.
    p_p6_q = find_p("Cara iterasi mana yang paling mudah menurut Anda? Jelaskan alasannya.")
    if p_p6_q:
        # Remove old text between p_p6_q and Praktikum 7
        curr = p_p6_q._p.getnext()
        while curr is not None:
            text = "".join(curr.itertext()).strip()
            if "Praktikum 7" in text:
                break
            nxt = curr.getnext()
            curr.getparent().remove(curr)
            curr = nxt

        cur_anchor = p_p6_q
        p_note = insert_callout_after(cur_anchor, doc,
            "Catatan Penting Alur Eksekusi Program:\n"
            "Karena pada Praktikum 5 sebelumnya telah dijalankan perintah mahasiswa.clear(), maka kondisi list mahasiswa saat ini adalah kosong ([] dengan length = 0). Oleh sebab itu, jika ketiga perulangan (for, for-in, dan forEach) langsung dijalankan, tidak ada satupun nama yang tercetak ke terminal karena kondisi perulangan langsung selesai seketika.\n\n"
            "[Langkah Tambahan untuk Demonstrasi]:\n"
            "Agar proses perulangan dapat diamati dan dibandingkan hasilnya, list diisi kembali terlebih dahulu dengan data contoh:\n"
            "mahasiswa.addAll(['Ahmad', 'Budi', 'Citra', 'Dewi']);",
            title="Analisis State Koleksi Setelah Praktikum 5",
            border_color="D97706", bg_color="FFFBEB"
        )

        p_iter_out = insert_p_after(p_note,
            "Keluaran Masing-Masing Metode Iterasi (Setelah Pengisian Kembali):\n"
            "Ahmad\nBudi\nCitra\nDewi",
            doc, style='CodeBlock', space_before=2, space_after=4
        )

        p_iter_ans = insert_p_after(p_iter_out,
            "Jawaban Cara Iterasi Terbaik dan Alasannya:\n"
            "Cara iterasi yang paling mudah, bersih, dan direkomendasikan adalah perulangan for-in (for (String nama in mahasiswa)).\n"
            "Alasan:\n"
            "1. Sintaks Ringkas dan Deklaratif: Kode sangat mudah dibaca layaknya kalimat bahasa manusia tanpa perlu mendeklarasikan variabel counter manual (seperti int i = 0) maupun mengelola kenaikan indeks (i++).\n"
            "2. Mencegah Kesalahan Batas Indeks (Off-by-One Error): Karena Dart menangani penelusuran elemen secara internal, pemrogram terlindungi dari bug batas indeks atau RangeError.\n"
            "3. Mendukung Kontrol Alur Penuh: Berbeda dengan method forEach() yang berbasis callback fungsi anonim (di mana statement return hanya melompati iterasi saat itu dan tidak bisa menggunakan break), for-in mendukung penggunaan break dan continue untuk mengendalikan alur perulangan secara fleksibel.",
            doc, style='Normal', space_before=2, space_after=6
        )

        # Insert Screenshot Terminal Praktikum 5-6
        insert_image_after(p_iter_ans, doc, get_img("terminal_praktikum_5_6.png"),
            "Gambar D.2: Bukti Eksekusi Terminal dart run bin/praktikum_dasar.dart (Praktikum 5–6)", width=Inches(5.4))
        print("Praktikum 6 explanation and terminal_praktikum_5_6.png inserted.")

    # --- PRAKTIKUM 8 ---
    # Prompt: jelaskan bahwa map() menghasilkan Iterable transformasi baru dan tidak langsung mengubah List sumber.
    # Jangan menyebut List sumber “imutabel”.
    p_p8_q = find_p("Jelaskan fungsi map():")
    if p_p8_q:
        curr = p_p8_q._p.getnext()
        while curr is not None:
            text = "".join(curr.itertext()).strip()
            if "Praktikum 9" in text:
                break
            nxt = curr.getnext()
            curr.getparent().remove(curr)
            curr = nxt

        cur_anchor = p_p8_q
        p_p8_out = insert_p_after(cur_anchor, "[AHMAD, BUDI, ANDI, CITRA, ANISA]", doc, style='CodeBlock', space_before=2, space_after=4)

        insert_p_after(p_p8_out,
            "Penjelasan Fungsi dan Karakteristik map():\n"
            "Method map() adalah higher-order method pada koleksi Dart yang berfungsi untuk mentransformasi/memetakan setiap elemen pada koleksi menjadi nilai atau bentuk data baru berdasarkan fungsi proyeksi yang didefinisikan.\n\n"
            "Karakteristik Kunci map():\n"
            "1. Menghasilkan Iterable Transformasi Baru: map() menghasilkan objek Iterable baru yang memuat hasil transformasi dan tidak langsung mengubah data pada List sumber aslinya. List awal tetap mempertahankan isinya semula.\n"
            "2. Evaluasi Tunda (Lazy Evaluation): map() menghasilkan objek MappedIterable di mana fungsi transformasi baru dieksekusi saat elemen tersebut diakses atau dievaluasi secara eksplisit (misalnya saat memanggil .toList()).\n"
            "3. Ukuran Koleksi Tetap (1-to-1 Mapping): Panjang koleksi hasil pemetaan selalu sama persis dengan panjang koleksi sumber aslinya (berbeda dengan where() yang dapat menyaring atau mengurangi jumlah elemen).",
            doc, style='Normal', space_before=2, space_after=6
        )
        print("Praktikum 8 explanation updated.")

    # --- PRAKTIKUM 9 ---
    p_p9_h = find_p("Praktikum 9")
    p_p9_code = None
    for p in doc.paragraphs:
        if "var descending = mahasiswa.reversed" in p.text:
            p_p9_code = p
            break

    if p_p9_code:
        p_p9_out = insert_p_after(p_p9_code, "[Ahmad, Andi, Anisa, Budi, Citra]\n[Citra, Budi, Anisa, Andi, Ahmad]", doc, style='CodeBlock', space_before=2, space_after=4)
        p_p9_expl = insert_p_after(p_p9_out,
            "Penjelasan Pengurutan Data (Sorting):\n"
            "• mahasiswa.sort(): Mengurutkan elemen list secara ascending (A ke Z) berdasarkan perbandingan urutan karakter leksikografis String. Operasi sort() bersifat in-place, artinya langsung memodifikasi susunan urutan elemen pada memori list tersebut.\n"
            "• mahasiswa.reversed.toList(): Properti getter .reversed menghasilkan Iterable dengan urutan elemen yang dibalik. Ketika dievaluasi dengan .toList(), terbentuk list baru yang berurutan descending (Z ke A) tanpa merusak susunan pada list mahasiswa awal.",
            doc, style='Normal', space_before=2, space_after=6
        )
        insert_image_after(p_p9_expl, doc, get_img("terminal_praktikum_7_9.png"),
            "Gambar D.3: Bukti Eksekusi Terminal dart run bin/praktikum_dasar.dart (Praktikum 7–9)", width=Inches(5.4))
        print("Praktikum 9 explanation and terminal_praktikum_7_9.png inserted.")

    # --- BAGIAN I: Latihan Eksplorasi Collection ---
    p_sec_i = find_p("I. Latihan Eksplorasi Collection")
    p_sec_j = find_p("J. Tugas Praktikum Individu")
    if p_sec_i and p_sec_j:
        # Find the code block for lulus = nilai.where...
        target_p = None
        for p in doc.paragraphs:
            if p.text.strip().startswith("var lulus = nilai.where"):
                target_p = p
                break
        if not target_p:
            target_p = p_sec_i

        # Insert Table of 8 Operations
        i_headers = ["No", "Operasi Koleksi", "Kode Dart", "Logika / Perhitungan", "Hasil"]
        i_rows = [
            ["1", "Hitung jumlah data", "nilai.length", "Menghitung total elemen di dalam List", "7"],
            ["2", "Cari nilai tertinggi", "nilai.reduce((c, n) => c > n ? c : n)", "Membandingkan elemen mencari nilai maksimum", "95"],
            ["3", "Cari nilai terendah", "nilai.reduce((c, n) => c < n ? c : n)", "Membandingkan elemen mencari nilai minimum", "65"],
            ["4", "Hitung nilai rata-rata", "nilai.reduce((a, b) => a + b) / nilai.length", "Total penjumlahan (560) dibagi jumlah data (7)", "80.0"],
            ["5", "Ambil data nilai >= 80", "nilai.where((n) => n >= 80).toList()", "Menyaring elemen bernilai lebih besar/sama dengan 80", "[80, 90, 85, 95]"],
            ["6", "Ambil data nilai < 80", "nilai.where((n) => n < 80).toList()", "Menyaring elemen bernilai lebih kecil dari 80", "[75, 65, 70]"],
            ["7", "Urutkan kecil ke besar", "(List.from(nilai)..sort())", "Pengurutan ascending (naik)", "[65, 70, 75, 80, 85, 90, 95]"],
            ["8", "Urutkan besar ke kecil", "(List.from(nilai)..sort((a, b) => b.compareTo(a)))", "Pengurutan descending (turun)", "[95, 90, 85, 80, 75, 70, 65]"]
        ]
        col_w = [Inches(0.4), Inches(1.5), Inches(1.8), Inches(1.8), Inches(1.2)]

        p_tbl_space = insert_table_after(target_p, doc, i_headers, i_rows, col_widths=col_w)

        p_code = insert_code_after(p_tbl_space,
            "// bin/eksplorasi_collection.da\n"
            "void main() {\n"
            "  List<int> nilai = [75, 80, 90, 65, 85, 95, 70];\n"
            "  print('1. Jumlah data: ${nilai.length}');\n"
            "  print('2. Nilai tertinggi: ${nilai.reduce((c, n) => c > n ? c : n)}');\n"
            "  print('3. Nilai terendah: ${nilai.reduce((c, n) => c < n ? c : n)}');\n"
            "  double avg = nilai.reduce((a, b) => a + b) / nilai.length;\n"
            "  print('4. Nilai rata-rata: ${avg.toStringAsFixed(1)}');\n"
            "  print('5. Nilai >= 80: ${nilai.where((n) => n >= 80).toList()}');\n"
            "  print('6. Nilai < 80: ${nilai.where((n) => n < 80).toList()}');\n"
            "  print('7. Urutan kecil ke besar: ${List.from(nilai)..sort()}');\n"
            "  print('8. Urutan besar ke kecil: ${List.from(nilai)..sort((a, b) => b.compareTo(a))}');\n"
            "}", doc
        )

        p_img1 = insert_image_after(p_code, doc, get_img("terminal_eksplorasi_bagian_i_n.png"),
            "Gambar I.1: Bukti Eksekusi Terminal dart run bin/eksplorasi_collection.dart (Bagian I dan Bagian N)", width=Inches(5.4))
        insert_image_after(p_img1, doc, get_img("app_eksplorasi_collection.png"),
            "Gambar I.2: Antarmuka Interaktif Modul Eksplorasi Collection pada Aplikasi Flutter Mobile", width=Inches(2.7))
        print("Bagian I completed with table, code, and screenshots.")

    # --- BAGIAN J & CHALLENGE: Aplikasi Mata Kuliah ---
    p_sec_j = find_p("J. Tugas Praktikum Individu")
    p_sec_k = find_p("K. Tugas Pengayaan")
    p_challenge = find_p("Challenge")
    if p_challenge:
        # Find paragraph after challenge explanation
        curr = p_challenge
        for p in doc.paragraphs:
            if p.text.strip().startswith("Tambahkan filter kategori SKS"):
                curr = p
                break

        # Insert detailed features explanation
        p_mk_desc = insert_p_after(curr,
            "Implementasi dan Penjelasan 10 Fitur Aplikasi Mata Kuliah:\n"
            "1. Menampilkan Daftar Mata Kuliah: Menggunakan ListView.builder dengan Card interaktif yang menampilkan kode mata kuliah, nama, bobot SKS, dan dosen pengampu.\n"
            "2. Menambahkan Mata Kuliah: Form dialog modal (AlertDialog) yang memvalidasi kelengkapan isian kode, nama, SKS, dan nama dosen.\n"
            "3. Mengubah Mata Kuliah: Tombol edit pada setiap kartu item memungkinkan pembaruan data secara aman dengan melacak identifier kode lama.\n"
            "4. Menghapus Mata Kuliah: Tombol hapus menghapus objek dari koleksi List lokal.\n"
            "5. Dialog Konfirmasi Hapus: AlertDialog konfirmasi ('Hapus data mata kuliah ini?') dengan tombol Batal dan Hapus untuk mencegah kehilangan data tak sengaja.\n"
            "6. Pencarian Real-Time: Input teks pencarian menyaring data secara dinamis berdasarkan kecocokan kode MK, nama MK, atau nama dosen.\n"
            "7. Menampilkan Jumlah Mata Kuliah: Menghitung total data aktif hasil filter secara reaktif menggunakan properti .length.\n"
            "8. Menghitung Total SKS: Menggunakan method fold<int>(0, (sum, mk) => sum + mk.sks) untuk menjumlahkan seluruh bobot SKS aktif.\n"
            "9. Sorting Nama Mata Kuliah A–Z: Mengurutkan daftar menggunakan method .sort((a, b) => a.nama.compareTo(b.nama)).\n"
            "10. Sorting Berdasarkan SKS: Mengurutkan mata kuliah berdasarkan bobot SKS; pada implementasi dipilih urutan naik (ascending: a.sks.compareTo(b.sks)).",
            doc, style='Normal', space_before=3, space_after=6
        )

        # Challenge filter metrics table
        j_headers = ["Pilihan Filter SKS", "Daftar Mata Kuliah yang Ditampilkan", "Jumlah MK", "Total SKS", "Rata-rata SKS"]
        j_rows = [
            ["Semua", "Pemrograman Mobile (3 SKS), Struktur Data & Algoritma (4 SKS),\nPemrograman Berorientasi Objek (3 SKS), Desain Antarmuka Pengguna (2 SKS)", "4", "12", "3.00"],
            ["2 SKS", "Desain Antarmuka Pengguna (2 SKS)", "1", "2", "2.00"],
            ["3 SKS", "Pemrograman Mobile (3 SKS), Pemrograman Berorientasi Objek (3 SKS)", "2", "6", "3.00"],
            ["4 SKS", "Struktur Data dan Algoritma (4 SKS)", "1", "4", "4.00"]
        ]
        col_w_j = [Inches(1.1), Inches(2.7), Inches(0.8), Inches(0.8), Inches(1.1)]
        p_tbl_j = insert_table_after(p_mk_desc, doc, j_headers, j_rows, col_widths=col_w_j)

        # Screenshots for Mata Kuliah
        p_sbs1 = insert_side_by_side_after(p_tbl_j, doc,
            get_img("app_matakuliah_daftar_ringkasan.png"), "Gambar J.1: Daftar MK & Ringkasan Metrik (Default)",
            get_img("app_matakuliah_search.png"), "Gambar J.2: Pencarian Mata Kuliah (Keyword 'Mobile')"
        )
        p_sbs2 = insert_side_by_side_after(p_sbs1, doc,
            get_img("app_matakuliah_tambah_dialog.png"), "Gambar J.3: Dialog Tambah Mata Kuliah",
            get_img("app_matakuliah_edit_dialog.png"), "Gambar J.4: Dialog Edit Mata Kuliah"
        )
        p_sbs3 = insert_side_by_side_after(p_sbs2, doc,
            get_img("app_matakuliah_delete_dialog.png"), "Gambar J.5: Dialog Konfirmasi Hapus Mata Kuliah",
            get_img("app_matakuliah_sort_sks.png"), "Gambar J.6: Pengurutan Berdasarkan SKS Naik (2 -> 3 -> 4)"
        )
        p_sbs4 = insert_side_by_side_after(p_sbs3, doc,
            get_img("app_matakuliah_filter_2sks.png"), "Gambar J.7: Filter 2 SKS (1 MK, Total 2 SKS, Rata-rata 2.00)",
            get_img("app_matakuliah_filter_3sks.png"), "Gambar J.8: Filter 3 SKS (2 MK, Total 6 SKS, Rata-rata 3.00)"
        )
        insert_image_after(p_sbs4, doc,
            get_img("app_matakuliah_filter_4sks.png"),
            "Gambar J.9: Filter 4 SKS (1 MK, Total 4 SKS, Rata-rata 4.00)", width=Inches(2.7)
        )
        print("Bagian J & Challenge completed with features, table, and screenshots.")

    # --- BAGIAN K: Tugas Pengayaan – Inventory Barang ---
    p_sec_k = find_p("K. Tugas Pengayaan")
    p_sec_l = find_p("L. Pertanyaan Analisis")
    if p_sec_k and p_sec_l:
        target_p = None
        for p in doc.paragraphs:
            if "Hitung nilai setiap barang" in p.text:
                target_p = p
                break
        if not target_p:
            target_p = p_sec_k

        k_headers = ["No", "Kode", "Nama Barang", "Stok", "Harga Satuan (Rp)", "Nilai Barang (Stok × Harga)"]
        k_rows = [
            ["1", "B001", "Laptop Asus ROG Zephyrus", "10", "Rp 15.000.000", "Rp 150.000.000"],
            ["2", "B002", "Mouse Wireless Logitech M330", "50", "Rp 250.000", "Rp 12.500.000"],
            ["3", "B003", "Keyboard Mechanical RGB", "30", "Rp 650.000", "Rp 19.500.000"],
            ["4", "B004", "Monitor LED 24 Inch IPS", "12", "Rp 2.200.000", "Rp 26.400.000"],
            ["5", "B005", "Flashdisk USB 3.0 64GB", "38", "Rp 50.000", "Rp 1.900.000"],
            ["TOTAL", "-", "Total 5 Barang", "140 unit", "-", "Rp 210.300.000"]
        ]
        col_w_k = [Inches(0.5), Inches(0.8), Inches(2.2), Inches(0.8), Inches(1.4), Inches(1.6)]
        p_tbl_k = insert_table_after(target_p, doc, k_headers, k_rows, col_widths=col_w_k)

        p_calc = insert_p_after(p_tbl_k,
            "Perhitungan Total Persediaan:\n"
            "Total Persediaan = Σ(stok × harga) = Rp 150.000.000 + Rp 12.500.000 + Rp 19.500.000 + Rp 26.400.000 + Rp 1.900.000 = Rp 210.300.000\n\n"
            "Implementasi Rumus pada Dart:\n"
            "double totalPersediaan = daftarBarang.fold(0.0, (sum, b) => sum + (b.stok * b.harga));",
            doc, style='Normal', space_before=2, space_after=6
        )

        p_sbs_k1 = insert_side_by_side_after(p_calc, doc,
            get_img("app_inventory_daftar_total.png"), "Gambar K.1: Daftar Barang & Total Persediaan Rp 210.300.000",
            get_img("app_inventory_search.png"), "Gambar K.2: Pencarian Barang (Keyword 'Logitech')"
        )
        p_sbs_k2 = insert_side_by_side_after(p_sbs_k1, doc,
            get_img("app_inventory_tambah_dialog.png"), "Gambar K.3: Dialog Tambah Data Barang",
            get_img("app_inventory_edit_dialog.png"), "Gambar K.4: Dialog Edit Data Barang"
        )
        insert_image_after(p_sbs_k2, doc,
            get_img("app_inventory_delete_dialog.png"),
            "Gambar K.5: Dialog Konfirmasi Hapus Data Barang", width=Inches(2.7)
        )
        print("Bagian K completed with table, calculation, and screenshots.")

    # --- BAGIAN L: Pertanyaan Analisis (10 Soal) ---
    p_sec_l = find_p("L. Pertanyaan Analisis")
    if p_sec_l:
        l_answers = {
            "1. Apa perbedaan List, Set, dan Map?": (
                "Jawaban:\n"
                "• List: Koleksi elemen berurutan (ordered) di mana setiap elemen diakses melalui indeks berbasis nol (0, 1, 2, ...). List mengizinkan adanya elemen duplikat (nilai yang sama).\n"
                "• Set: Koleksi sekumpulan elemen yang setiap nilainya harus unik (tidak mengizinkan duplikat). Operasi penambahan elemen yang sudah ada akan diabaikan. Set tidak menjamin urutan posisi elemen.\n"
                "• Map: Koleksi berbasis pasangan kunci dan nilai (key-value pair). Setiap key bersifat unik dan bertindak sebagai identifier/pengindeks untuk mengakses value pasangannya."
            ),
            "2. Mengapa indeks List dimulai dari 0?": (
                "Jawaban:\n"
                "Indeks List dimulai dari 0 (zero-based indexing) karena secara historis berakar dari arsitektur komputer dan bahasa pemrograman tingkat rendah (seperti C). Indeks merepresentasikan nilai jarak perpindahan (offset/jarak pergeseran) dari alamat memori awal (base address) array. Elemen pertama berada tepat di alamat awal sehingga jarak offset-nya adalah 0 × ukuran elemen = 0. Elemen kedua bergeser 1 unit ukuran elemen (offset 1), dan seterusnya."
            ),
            "3. Apa fungsi add()?": (
                "Jawaban:\n"
                "Method add() berfungsi untuk menambahkan satu elemen baru ke posisi paling akhir dari List. Operasi ini memperbesar ukuran (length) List secara dinamis sebanyak satu elemen."
            ),
            "4. Apa perbedaan remove() dan removeAt()?": (
                "Jawaban:\n"
                "• remove(value): Menghapus elemen berdasarkan kecocokan nilai objek (equality). Method ini mencari elemen bernilai sama dari awal list, menghapus kemunculan pertamanya, dan mengembalikan boolean (true jika berhasil ditemukan dan dihapus, false jika tidak ditemukan).\n"
                "• removeAt(index): Menghapus elemen berdasarkan posisi nomor indeks numeriknya di dalam List. Method ini mengembalikan objek elemen yang dihapus dan menggeser posisi seluruh elemen berikutnya ke kiri. Jika indeks berada di luar rentang, method ini melempar exception RangeError."
            ),
            "5. Apa fungsi where()?": (
                "Jawaban:\n"
                "Method where() adalah higher-order method yang berfungsi untuk menyaring (memfilter) elemen-elemen di dalam koleksi berdasarkan fungsi kondisi (predicate boolean). Method ini mengembalikan Iterable baru yang hanya memuat elemen-elemen yang menghasilkan nilai true pada fungsi pengujian tersebut."
            ),
            "6. Apa fungsi map()?": (
                "Jawaban:\n"
                "Method map() adalah higher-order method yang berfungsi untuk mentransformasikan/memetakan setiap elemen pada koleksi menjadi nilai atau bentuk data baru berdasarkan formula fungsi proyeksi yang diberikan, tanpa langsung mengubah data pada List sumber aslinya (menghasilkan Iterable baru)."
            ),
            "7. Apa kegunaan toList()?": (
                "Jawaban:\n"
                "Method toList() digunakan untuk mengevaluasi dan mengonversi objek Iterable (seperti hasil operasi where() atau map() yang bersifat lazy evaluation / evaluasi tunda) menjadi objek List konkret di memori, sehingga elemen-elemennya dapat diakses menggunakan indeks [], diukur panjangnya (.length), serta dimanipulasi dengan method List lainnya."
            ),
            "8. Mengapa aplikasi Flutter menggunakan setState() setelah data List berubah?": (
                "Jawaban:\n"
                "Dalam arsitektur StatefulWidget Flutter, pemanggilan setState() memberitahu framework bahwa state internal objek telah mengalami perubahan. Framework Flutter kemudian akan menandai widget tersebut sebagai 'dirty' dan menjadwalkan eksekusi ulang method build() pada frame berikutnya, sehingga antarmuka visual (UI) langsung dirender ulang sesuai data koleksi terbaru."
            ),
            "9. Apa keuntungan menggunakan ListView.builder dibanding menuliskan widget secara manual?": (
                "Jawaban:\n"
                "ListView.builder menerapkan konsep 'lazy rendering' dan 'viewport recycling', yaitu hanya merender dan membuat widget item yang sedang berada di layar tampak (viewport). Item yang digeser keluar layar akan didaur ulang di memori. Hal ini menghemat alokasi memori RAM dan penggunaan CPU secara drastis dibandingkan menulis widget manual (misal SingleChildScrollView + Column) yang merender seluruh data sekaligus sejak awal."
            ),
            "10. Mengapa model Mahasiswa lebih baik dibanding menggunakan banyak List<String> terpisah?": (
                "Jawaban:\n"
                "Menggunakan class model Mahasiswa menerapkan prinsip Enkapsulasi Berorientasi Objek (OOP). Seluruh atribut yang saling terkait (NIM, nama, prodi, semester, IPK) disatukan dalam satu kesatuan objek tunggal. Hal ini menjamin integritas data (mencegah bug desinkronisasi indeks saat operasi tambah, hapus, atau sorting), memberikan keamanan tipe data (type-safety) saat kompilasi (misalnya IPK bertipe double, bukan String), serta memudahkan pemeliharaan kode (clean code)."
            )
        }

        for q_prefix, ans_text in l_answers.items():
            p_q = find_p(q_prefix[:15])
            if p_q:
                insert_p_after(p_q, ans_text, doc, style='Normal', space_before=2, space_after=6)
        print("Bagian L 10 analysis questions completed.")

    # --- BAGIAN M: Analisis Kesalahan Program ---
    p_sec_m = find_p("M. Analisis Kesalahan Program")
    if p_sec_m:
        p_m_q1 = find_p("Apa yang terjadi?")
        if p_m_q1:
            insert_p_after(p_m_q1,
                "Jawaban:\n"
                "Program mengalami kegagalan fatal saat dijalankan dan melempar runtime exception:\n"
                "RangeError (Index out of range: index should be less than 2: index 5)\n"
                "Eksekusi program berhenti seketika (crash) karena sistem menolak pengaksesan memori pada indeks yang tidak teralokasi.",
                doc, style='Normal', space_before=2, space_after=6
            )

        p_m_q2 = find_p("Mengapa hal tersebut terjadi?")
        if p_m_q2:
            insert_p_after(p_m_q2,
                "Jawaban:\n"
                "List mahasiswa hanya dideklarasikan dengan 2 elemen string (['Ahmad', 'Budi']). Karena Dart menggunakan sistem indeks berbasis 0, indeks valid yang tersedia hanyalah 0 dan 1 (rentang 0 sampai length - 1, di mana length = 2). Ketika program mengeksekusi mahasiswa[5], program mencoba mengakses indeks ke-5 yang berada jauh di luar batas kapasitas list tersebut.",
                doc, style='Normal', space_before=2, space_after=6
            )

        p_m_q3 = find_p("Bagaimana cara menghindarinya?")
        if p_m_q3:
            insert_p_after(p_m_q3,
                "Jawaban:\n"
                "Empat cara utama untuk menghindari RangeError:\n"
                "1. Validasi Batas Indeks: Selalu periksa apakah indeks berada dalam rentang valid sebelum mengakses:\n"
                "   if (index >= 0 && index < mahasiswa.length) { print(mahasiswa[index]); }\n"
                "2. Gunakan Getter atau Ekstensi Aman: Memanfaatkan getter bawaan seperti .firstOrNull atau method elementAtOrNull(index) (dari package:collection):\n"
                "   print(mahasiswa.elementAtOrNull(5)); // Mengembalikan null, bukan crash\n"
                "3. Gunakan Perulangan Aman (for-in / forEach): Hindari perulangan indeks manual dengan beralih ke iterasi otomatis yang dikontrol koleksi:\n"
                "   for (var m in mahasiswa) { print(m); }\n"
                "4. Penanganan Exception (Error Handling): Membungkus pemanggilan berisiko dengan blok try-catch:\n"
                "   try { print(mahasiswa[5]); } catch (e) { print('Indeks tidak valid: $e'); }",
                doc, style='Normal', space_before=2, space_after=6
            )
        print("Bagian M completed.")

    # --- BAGIAN N: Problem Solving ---
    p_sec_n = find_p("N. Problem Solving")
    if p_sec_n:
        target_p = None
        for p in doc.paragraphs:
            if "Tuliskan jawaban di sini" in p.text:
                target_p = p
                break
        if not target_p:
            target_p = p_sec_n

        p_code_n = insert_code_after(target_p,
            "void main() {\n"
            "  List<int> nilai = [80, 65, 90, 70, 95];\n"
            "\n"
            "  // Mengambil nilai >= 80 menggunakan where()\n"
            "  var hasil = nilai.where((n) => n >= 80).toList();\n"
            "\n"
            "  // Menghitung jumlah data yang memenuhi kondisi\n"
            "  int jumlah = hasil.length;\n"
            "\n"
            "  print('Nilai >= 80 : $hasil');\n"
            "  print('Jumlah data : $jumlah');\n"
            "}", doc
        )

        insert_p_after(p_code_n,
            "Hasil Keluaran Program:\n"
            "Nilai >= 80 : [80, 90, 95]\n"
            "Jumlah data : 3\n\n"
            "Penjelasan Singkat:\n"
            "Method where((n) => n >= 80) menyaring elemen yang bernilai lebih besar atau sama dengan 80, menghasilkan elemen 80, 90, dan 95. Method .toList() mengonversi hasil penyaringan menjadi List<int> baru, dan properti .length menghitung total data yang memenuhi kriteria, yaitu tepat sebanyak 3 data.",
            doc, style='Normal', space_before=2, space_after=6
        )
        print("Bagian N completed with solution and output.")

    # --- BAGIAN O & Q: Mini Project – Student Management App ---
    p_sec_o = find_p("O. Mini Project")
    p_sec_p = find_p("P. Algoritma Mini Project")
    if p_sec_o and p_sec_p:
        # Insert Mini Project Features & Screenshots
        target_p = None
        for p in doc.paragraphs:
            if "Dashboard minimal menampilkan" in p.text:
                target_p = p
                break
        if not target_p:
            target_p = p_sec_o

        p_o_desc = insert_p_after(target_p,
            "Rincian Fitur Student Management App yang Berhasil Dibangun:\n"
            "1. Dashboard Ringkasan Metrik: Menampilkan kartu metrik dinamis: Total Mahasiswa Aktif, Rata-rata IPK, dan Mahasiswa Berprestasi (Cumlaude / IPK >= 3.50).\n"
            "2. Manajemen Data (CRUD Lengkap):\n"
            "   • CREATE: Dialog penambahan mahasiswa baru dengan validasi form (NIM, Nama, Prodi, Semester, IPK).\n"
            "   • READ: Tampilan daftar kartu mahasiswa menggunakan ListView.builder interaktif.\n"
            "   • UPDATE: Dialog edit data mahasiswa yang mendukung pembaruan seluruh field termasuk identitas lama.\n"
            "   • DELETE: Tombol hapus yang dilengkapi dialog konfirmasi interaktif.\n"
            "3. Pencarian Dinamis: Menyaring data mahasiswa secara real-time berdasarkan kecocokan NIM, Nama, atau Program Studi.\n"
            "4. Filter Kategori IPK: Opsi filter 'Semua', 'IPK >= 3.50' (Cumlaude), dan 'IPK < 3.50'.\n"
            "5. Lima Pilihan Sorting: Pengurutan Nama A-Z, Nama Z-A, IPK Tertinggi, IPK Terendah, dan Semester.",
            doc, style='Normal', space_before=2, space_after=6
        )

        # Screenshots pairs for Mini Project
        p_sbs_o1 = insert_side_by_side_after(p_o_desc, doc,
            get_img("app_mahasiswa_dashboard_daftar.png"), "Gambar O.1: Dashboard Metrik & Daftar Mahasiswa (Default)",
            get_img("app_mahasiswa_search.png"), "Gambar O.2: Pencarian Mahasiswa (Keyword 'Citra')"
        )
        p_sbs_o2 = insert_side_by_side_after(p_sbs_o1, doc,
            get_img("app_mahasiswa_tambah_dialog.png"), "Gambar O.3: Dialog Tambah Mahasiswa",
            get_img("app_mahasiswa_tambah_hasil.png"), "Gambar O.4: Hasil Penambahan Mahasiswa Baru (Eka Pratama)"
        )
        p_sbs_o3 = insert_side_by_side_after(p_sbs_o2, doc,
            get_img("app_mahasiswa_edit_dialog.png"), "Gambar O.5: Dialog Edit Data Mahasiswa",
            get_img("app_mahasiswa_edit_hasil.png"), "Gambar O.6: Hasil Edit Data Mahasiswa (Ahmad Fauzi Updated)"
        )
        p_sbs_o4 = insert_side_by_side_after(p_sbs_o3, doc,
            get_img("app_mahasiswa_delete_dialog.png"), "Gambar O.7: Dialog Konfirmasi Hapus Mahasiswa",
            get_img("app_mahasiswa_delete_hasil.png"), "Gambar O.8: Hasil Setelah Mahasiswa Dihapus (Sisa 3 Data)"
        )
        p_sbs_o5 = insert_side_by_side_after(p_sbs_o4, doc,
            get_img("app_mahasiswa_filter_cumlaude.png"), "Gambar O.9: Filter IPK >= 3.50 (2 Mahasiswa Cumlaude)",
            get_img("app_mahasiswa_filter_non_cumlaude.png"), "Gambar O.10: Filter IPK < 3.50 (2 Mahasiswa Non-Cumlaude)"
        )
        p_sbs_o6 = insert_side_by_side_after(p_sbs_o5, doc,
            get_img("app_mahasiswa_sort_nama_za.png"), "Gambar O.11: Sorting Nama Z ke A (Citra -> Budi -> Ahmad)",
            get_img("app_mahasiswa_sort_ipk_tertinggi.png"), "Gambar O.12: Sorting IPK Tertinggi (3.75 -> 3.60 -> 3.45)"
        )
        insert_side_by_side_after(p_sbs_o6, doc,
            get_img("app_mahasiswa_sort_ipk_terendah.png"), "Gambar O.13: Sorting IPK Terendah (3.20 -> 3.45 -> ...)",
            get_img("app_mahasiswa_sort_semester.png"), "Gambar O.14: Sorting Berdasarkan Semester"
        )
        print("Bagian O completed with features and 7 screenshot pairs.")

    # --- BAGIAN Q: Indikator Keberhasilan Praktikum ---
    p_sec_q = find_p("Q. Indikator Keberhasilan")
    p_sec_r = find_p("R. Rubrik Penilaian")
    if p_sec_q and p_sec_r:
        # Find the last paragraph before R
        target_p = None
        for p in doc.paragraphs:
            if "10. Mahasiswa dapat menjelaskan" in p.text:
                target_p = p
                break
        if not target_p:
            target_p = p_sec_q

        q_headers = ["No", "Indikator Keberhasilan", "Kriteria Penilaian", "Status", "Bukti Verifikasi / Pengujian"]
        q_rows = [
            ["1", "Project Flutter dapat dijalankan tanpa error", "Aplikasi dapat di-build dan berjalan lancar tanpa exception", "TERPENUHI", "Verifikasi terminal flutter run -d chrome & flutter build web sukses"],
            ["2", "Data dapat ditampilkan dari List", "Data objek model berhasil dirender ke kartu antarmuka", "TERPENUHI", "Tampilan kartu data mahasiswa dari List lokal (Gambar O.1)"],
            ["3", "Data baru dapat ditambahkan", "Input form dialog menambah objek baru ke dalam List", "TERPENUHI", "Pengisian form dialog tambah dan data tampil di UI (Gambar O.3 & O.4)"],
            ["4", "Data dapat diedit", "Perubahan atribut objek lama terbarui secara konsisten", "TERPENUHI", "Dialog edit memperbarui data lama secara instan (Gambar O.5 & O.6)"],
            ["5", "Data dapat dihapus", "Menghapus elemen dari List dengan konfirmasi aman", "TERPENUHI", "Dialog konfirmasi hapus dan data terhapus dari UI (Gambar O.7 & O.8)"],
            ["6", "Search berjalan dengan benar", "Pencarian real-time menyaring data berdasarkan keyword", "TERPENUHI", "Pencarian nama, NIM, atau prodi dengan where() (Gambar O.2)"],
            ["7", "Sorting berjalan dengan benar", "Data dapat diurutkan berdasarkan berbagai kriteria", "TERPENUHI", "5 mode sorting berjalan sesuai ekspektasi (Gambar O.11 - O.14)"],
            ["8", "ListView.builder digunakan", "Menerapkan lazy rendering untuk efisiensi memori", "TERPENUHI", "Implementasi ListView.builder pada lib/pages/mahasiswa_page.dart"],
            ["9", "Perubahan data langsung terlihat pada UI", "State reaktif memperbarui visual antarmuka seketika", "TERPENUHI", "Pemanggilan setState() pada seluruh mutasi data CRUD"],
            ["10", "Penjelasan Collection, State, dan Widget", "Pemahaman mendalam tentang hubungan konsep inti", "TERPENUHI", "Penjelasan komprehensif tertulis pada Bagian L, T, dan V"]
        ]
        col_w_q = [Inches(0.4), Inches(2.2), Inches(2.3), Inches(1.1), Inches(2.2)]
        p_tbl_q = insert_table_after(target_p, doc, q_headers, q_rows, col_widths=col_w_q)

        # Insert Terminal Verification Screenshots
        p_t_an = insert_image_after(p_tbl_q, doc, get_img("terminal_flutter_analyze.png"),
            "Gambar Q.1: Bukti Eksekusi Terminal flutter analyze (Bebas Masalah Sintaks dan Linter / No Issues Found)", width=Inches(5.4))
        insert_image_after(p_t_an, doc, get_img("terminal_flutter_test.png"),
            "Gambar Q.2: Bukti Eksekusi Terminal flutter test (Seluruh 28 Unit & Widget Test Berhasil Lolos / 28 Passed)", width=Inches(5.4))
        print("Bagian Q completed with checklist table and terminal verification screenshots.")

    # --- BAGIAN U: Refleksi Mahasiswa ---
    p_sec_u = find_p("U. Refleksi Mahasiswa")
    if p_sec_u:
        u_answers = {
            "Konsep baru yang saya pahami:": (
                "[Catatan: Draf ini disusun berdasarkan pengerjaan proyek riil. Silakan sesuaikan dengan pengalaman pribadi Anda]\n"
                "Saya memahami secara mendalam cara kerja struktur data collection pada Dart (List, Set, dan Map), khususnya higher-order methods seperti where() untuk memfilter elemen, map() untuk transformasi data, dan sort() untuk pengurutan. Selain itu, saya memahami keterkaitan erat antara manipulasi collection di memori dengan siklus hidup state reaktif Flutter (setState()) dan optimasi perenderan antarmuka dinamis menggunakan ListView.builder."
            ),
            "Bagian praktikum yang paling sulit:": (
                "[Catatan: Draf ini disusun berdasarkan pengerjaan proyek riil. Silakan sesuaikan dengan pengalaman pribadi Anda]\n"
                "Bagian yang paling menantang adalah menjaga integritas dan konsistensi data saat melakukan operasi UPDATE pada entitas objek. Khususnya ketika primary identifier (seperti NIM atau Kode Mata Kuliah) diperbolehkan diedit oleh pengguna, repository harus melacak identitas data lama agar tidak terjadi duplikasi data atau kesalahan penimpaan indeks. Selain itu, sinkronisasi metrik statistik dashboard (rata-rata IPK dan total SKS) secara real-time saat filter diterapkan membutuhkan logika perhitungan yang teliti."
            ),
            "Solusi yang saya lakukan:": (
                "[Catatan: Draf ini disusun berdasarkan pengerjaan proyek riil. Silakan sesuaikan dengan pengalaman pribadi Anda]\n"
                "Saya memisahkan kode logika bisnis dan penyimpanan data ke dalam lapisan in-memory repository (MahasiswaRepository, MataKuliahRepository, dan InventoryRepository) yang memiliki method update berbasis targetNim / targetKode. Selain itu, saya menyusun unit testing dan widget testing secara komprehensif (28 test case) untuk memastikan seluruh operasi CRUD, filter, dan sorting terverifikasi bebas bug sebelum dijalankan pada UI."
            ),
            "Fitur tambahan yang berhasil saya buat:": (
                "[Catatan: Draf ini disusun berdasarkan pengerjaan proyek riil. Silakan sesuaikan dengan pengalaman pribadi Anda]\n"
                "Saya berhasil mengembangkan antarmuka terpadu berbasis multi-tab (NavigationRail dan bottom navigation) yang mencakup 4 modul lengkap:\n"
                "1. Modul Mahasiswa (dengan metrik dashboard IPK, filter cumlaude, dan 5 mode sorting).\n"
                "2. Modul Mata Kuliah (dengan perhitungan dinamis total & rata-rata SKS serta filter 2, 3, 4 SKS).\n"
                "3. Modul Inventory Barang (dengan perhitungan nilai stok per barang dan total nilai persediaan otomatis).\n"
                "4. Modul Eksplorasi GUI Koleksi Dart (antarmuka interaktif visual untuk menguji 8 operasi koleksi)."
            ),
            "Kesimpulan pembelajaran hari ini:": (
                "[Catatan: Draf ini disusun berdasarkan pengerjaan proyek riil. Silakan sesuaikan dengan pengalaman pribadi Anda]\n"
                "Penguasaan collection (khususnya List<Model>) adalah keterampilan fundamental terpenting dalam pemrograman aplikasi mobile. Struktur data ini berfungsi sebagai jembatan representasi state lokal sebelum aplikasi diintegrasikan dengan sumber data eksternal permanen seperti database SQLite lokal maupun REST API backend."
            )
        }

        for q_prefix, ans_text in u_answers.items():
            p_q = find_p(q_prefix[:20])
            if p_q:
                insert_p_after(p_q, ans_text, doc, style='Normal', space_before=2, space_after=6)
        print("Bagian U reflection completed.")

    # --- BAGIAN V: Tugas Pertemuan Berikutnya ---
    p_sec_v = find_p("V. Tugas Pertemuan Berikutnya")
    if p_sec_v:
        curr = p_sec_v
        for p in doc.paragraphs:
            if "Kembangkan aplikasi dari List lokal" in p.text:
                curr = p
                break

        insert_p_after(curr,
            "Penjelasan Arsitektur Pemisahan Model, Repository, dan UI pada Proyek Ini:\n"
            "Proyek ini telah mengimplementasikan prinsip pemisahan tanggung jawab (Separation of Concerns) yang membagi kode ke dalam 3 lapisan arsitektural utama:\n\n"
            "1. Layer Model (lib/models/):\n"
            "   Memuat kelas murni Mahasiswa, MataKuliah, dan Barang. Lapisan ini murni merepresentasikan entitas data dengan atribut spesifik, validasi tipe data (type safety), constructor berparameter, dan method salinan (copyWith).\n\n"
            "2. Layer Repository / In-Memory Storage (lib/repositories/):\n"
            "   Memuat MahasiswaRepository, MataKuliahRepository, dan InventoryRepository. Lapisan ini mengisolasi seluruh logika manipulasi List lokal (tambah, baca, ubah, hapus, cari, saring, dan urutkan). Komponen UI tidak pernah menyentuh List mentah secara langsung, melainkan selalu berinteraksi melalui antarmuka repository ini.\n\n"
            "3. Layer Presentation / UI (lib/pages/):\n"
            "   Memuat HomePage, MahasiswaPage, MataKuliahPage, InventoryPage, dan EksplorasiCollectionPage. Lapisan ini hanya berfokus pada rendering widget visual, penangkapan interaksi pengguna (touch/click), pemanggilan method repository, dan pembaruan state visual menggunakan setState().\n\n"
            "Roadmap Transisi Menuju Database Lokal dan REST API Backend:\n"
            "Karena antarmuka data telah diabstraksikan di lapisan repository, transisi arsitektur pada pertemuan berikutnya dapat dilakukan dengan sangat elegan:\n"
            "• Transisi ke SQLite (Database Lokal): Cukup membuat implementasi baru SqliteMahasiswaRepository menggunakan package sqflite yang mengimplementasikan operasi SQL (INSERT, SELECT, UPDATE, DELETE). Komponen UI tidak perlu diubah sama sekali.\n"
            "• Transisi ke REST API Backend: Cukup membuat ApiMahasiswaRepository yang mengirimkan HTTP request (GET, POST, PUT, DELETE) menggunakan package http atau dio. UI tetap mengonsumsi data dengan kontrak pemanggilan yang sama.",
            doc, style='Normal', space_before=2, space_after=8
        )
        print("Bagian V architectural explanation added.")

    # --- LAMPIRAN: DAFTAR BUKTI EKSEKUSI SCREENSHOT ---
    p_ref = find_p("Referensi Singkat")
    if p_ref:
        p_last = doc.paragraphs[-1]
        p_lampiran_h = insert_p_after(p_last, "Lampiran: Matriks Bukti Eksekusi dan Tangkapan Layar Asli", doc, style='Heading 1', space_before=14, space_after=6)
        p_lampiran_desc = insert_p_after(p_lampiran_h,
            "Seluruh tangkapan layar dalam dokumen ini merupakan bukti eksekusi nyata (bukan mockup buatan) yang diambil langsung dari terminal CLI dan aplikasi Flutter berjalan pada viewport ponsel 390×844 px. Berkas gambar asli tersimpan pada folder bukti_screenshot/ dengan pemetaan sebagai berikut:",
            doc, style='Normal', space_before=2, space_after=6
        )

        lamp_headers = ["No", "Kategori / Modul", "Bagian LKM", "Nama Berkas Screenshot", "Fitur / Pembuktian"]
        lamp_rows = [
            ["1", "Terminal CLI", "D (Praktikum 1–4)", "terminal_praktikum_1_4.png", "Eksekusi create, read, add, update List"],
            ["2", "Terminal CLI", "D (Praktikum 5–6)", "terminal_praktikum_5_6.png", "Eksekusi remove, clear, dan for/for-in/forEach"],
            ["3", "Terminal CLI", "D (Praktikum 7–9)", "terminal_praktikum_7_9.png", "Eksekusi where, map, dan sort/reversed"],
            ["4", "Terminal CLI", "Bagian I & N", "terminal_eksplorasi_bagian_i_n.png", "8 operasi eksplorasi nilai & problem solving"],
            ["5", "Flutter Mobile", "Bagian I", "app_eksplorasi_collection.png", "Antarmuka GUI interaktif koleksi Dart"],
            ["6", "Flutter Mobile", "Bagian J & Challenge", "app_matakuliah_daftar_ringkasan.png", "Daftar MK & ringkasan metrik SKS"],
            ["7", "Flutter Mobile", "Bagian J", "app_matakuliah_search.png", "Pencarian kata kunci 'Mobile'"],
            ["8", "Flutter Mobile", "Bagian J", "app_matakuliah_tambah_dialog.png", "Form dialog tambah mata kuliah baru"],
            ["9", "Flutter Mobile", "Bagian J", "app_matakuliah_edit_dialog.png", "Form dialog edit mata kuliah"],
            ["10", "Flutter Mobile", "Bagian J", "app_matakuliah_delete_dialog.png", "Dialog konfirmasi hapus mata kuliah"],
            ["11", "Flutter Mobile", "Bagian J", "app_matakuliah_sort_sks.png", "Pengurutan berdasarkan SKS naik"],
            ["12", "Flutter Mobile", "Challenge J", "app_matakuliah_filter_2sks.png", "Filter 2 SKS (1 MK, Total 2 SKS)"],
            ["13", "Flutter Mobile", "Challenge J", "app_matakuliah_filter_3sks.png", "Filter 3 SKS (2 MK, Total 6 SKS)"],
            ["14", "Flutter Mobile", "Challenge J", "app_matakuliah_filter_4sks.png", "Filter 4 SKS (1 MK, Total 4 SKS)"],
            ["15", "Flutter Mobile", "Bagian K", "app_inventory_daftar_total.png", "Daftar barang & total persediaan Rp 210.300.000"],
            ["16", "Flutter Mobile", "Bagian K", "app_inventory_search.png", "Pencarian barang kata kunci 'Logitech'"],
            ["17", "Flutter Mobile", "Bagian K", "app_inventory_tambah_dialog.png", "Form dialog tambah barang baru"],
            ["18", "Flutter Mobile", "Bagian K", "app_inventory_edit_dialog.png", "Form dialog edit stok/harga barang"],
            ["19", "Flutter Mobile", "Bagian K", "app_inventory_delete_dialog.png", "Dialog konfirmasi hapus barang"],
            ["20", "Flutter Mobile", "Bagian O", "app_mahasiswa_dashboard_daftar.png", "Dashboard 3 metrik & ListView.builder"],
            ["21", "Flutter Mobile", "Bagian O", "app_mahasiswa_search.png", "Pencarian mahasiswa kata kunci 'Citra'"],
            ["22", "Flutter Mobile", "Bagian O", "app_mahasiswa_tambah_dialog.png", "Form dialog tambah mahasiswa baru"],
            ["23", "Flutter Mobile", "Bagian O", "app_mahasiswa_tambah_hasil.png", "Hasil penambahan mahasiswa di daftar"],
            ["24", "Flutter Mobile", "Bagian O", "app_mahasiswa_edit_dialog.png", "Form dialog edit data mahasiswa"],
            ["25", "Flutter Mobile", "Bagian O", "app_mahasiswa_edit_hasil.png", "Hasil edit data mahasiswa terbarui"],
            ["26", "Flutter Mobile", "Bagian O", "app_mahasiswa_delete_dialog.png", "Dialog konfirmasi hapus data mahasiswa"],
            ["27", "Flutter Mobile", "Bagian O", "app_mahasiswa_delete_hasil.png", "Hasil setelah data mahasiswa dihapus"],
            ["28", "Flutter Mobile", "Bagian O", "app_mahasiswa_filter_cumlaude.png", "Filter IPK >= 3.50 (Mahasiswa Cumlaude)"],
            ["29", "Flutter Mobile", "Bagian O", "app_mahasiswa_filter_non_cumlaude.png", "Filter IPK < 3.50 (Mahasiswa Non-Cumlaude)"],
            ["30", "Flutter Mobile", "Bagian O", "app_mahasiswa_sort_nama_za.png", "Sorting nama descending (Z ke A)"],
            ["31", "Flutter Mobile", "Bagian O", "app_mahasiswa_sort_ipk_tertinggi.png", "Sorting IPK tertinggi ke terendah"],
            ["32", "Flutter Mobile", "Bagian O", "app_mahasiswa_sort_ipk_terendah.png", "Sorting IPK terendah ke tertinggi"],
            ["33", "Flutter Mobile", "Bagian O", "app_mahasiswa_sort_semester.png", "Sorting berdasarkan semester"],
            ["34", "Terminal CLI", "Bagian Q", "terminal_flutter_analyze.png", "Verifikasi flutter analyze (0 issues found)"],
            ["35", "Terminal CLI", "Bagian Q", "terminal_flutter_test.png", "Verifikasi flutter test (28 passed)"]
        ]
        col_w_lamp = [Inches(0.4), Inches(1.2), Inches(1.3), Inches(2.2), Inches(2.3)]
        insert_table_after(p_lampiran_desc, doc, lamp_headers, lamp_rows, col_widths=col_w_lamp)
        print("Lampiran matrix table added.")

    # --- SAVE DOCUMENT ---
    print(f"Saving final document to: {OUTPUT_DOCX}")
    doc.save(OUTPUT_DOCX)
    print("SUCCESS: LKM_Mengelola_Data_List_Koleksi_Flutter_Dart_LENGKAP.docx has been created!")
    print(f"Final file size: {os.path.getsize(OUTPUT_DOCX):,} bytes")

if __name__ == "__main__":
    main()
