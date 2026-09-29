import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def remove_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_callout_box(doc, text, title=None, border_color="2A4D69", bg_color="F5F7FA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)

    # Border & Shading
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')

    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24') # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_color)
    tcBorders.append(left)

    for b_name in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), bg_color)
    tcPr.append(shd)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    if title:
        r_t = p.add_run(f"{title}\n")
        r_t.font.bold = True
        r_t.font.size = Pt(10)
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    return tbl

def add_single_image(doc, img_path, caption, width=Inches(5.6)):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(img_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    r_cap.font.italic = True
    r_cap.font.size = Pt(9)
    r_cap.font.color.rgb = RGBColor(80, 80, 80)
    return p

def add_side_by_side_images(doc, img1_path, cap1, img2_path, cap2, img_width=Inches(2.7)):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_borders(tbl)

    # Cell 1
    c1 = tbl.cell(0, 0)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(6)
    p1.paragraph_format.space_after = Pt(2)
    if os.path.exists(img1_path):
        p1.add_run().add_picture(img1_path, width=img_width)
    cap_p1 = c1.add_paragraph()
    cap_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p1.paragraph_format.space_before = Pt(2)
    cap_p1.paragraph_format.space_after = Pt(6)
    r1 = cap_p1.add_run(cap1)
    r1.font.italic = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = RGBColor(80, 80, 80)

    # Cell 2
    c2 = tbl.cell(0, 1)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(2)
    if os.path.exists(img2_path):
        p2.add_run().add_picture(img2_path, width=img_width)
    cap_p2 = c2.add_paragraph()
    cap_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p2.paragraph_format.space_before = Pt(2)
    cap_p2.paragraph_format.space_after = Pt(6)
    r2 = cap_p2.add_run(cap2)
    r2.font.italic = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(80, 80, 80)

    # Add spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(6)
    return tbl

print("Helper definitions loaded successfully.")
